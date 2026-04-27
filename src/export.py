"""엑셀 다운로드 (조건부 서식 포함)."""

from __future__ import annotations

import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.formatting import (
    BLUE_BG,
    BLUE_TEXT,
    GRAY_BG,
    GRAY_TEXT,
    PERCENT_COLUMNS,
    PINK_BG,
    PINK_TEXT,
    is_money_column,
)


PINK_FILL = PatternFill(start_color=PINK_BG.lstrip("#"), end_color=PINK_BG.lstrip("#"), fill_type="solid")
BLUE_FILL = PatternFill(start_color=BLUE_BG.lstrip("#"), end_color=BLUE_BG.lstrip("#"), fill_type="solid")
GRAY_FILL = PatternFill(start_color=GRAY_BG.lstrip("#"), end_color=GRAY_BG.lstrip("#"), fill_type="solid")
HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")

PINK_FONT = Font(color=PINK_TEXT.lstrip("#"), bold=True)
BLUE_FONT = Font(color=BLUE_TEXT.lstrip("#"), bold=True)
GRAY_FONT = Font(color=GRAY_TEXT.lstrip("#"))
HEADER_FONT = Font(color="FFFFFF", bold=True)


def export_to_excel(
    main_df: pd.DataFrame,
    unmatched_df: pd.DataFrame | None = None,
    kpi: dict | None = None,
    period_label: str = "",
) -> bytes:
    """모든 데이터를 하나의 엑셀로.

    시트 구성:
    - 요약: KPI + 기간 정보
    - 광고별상세: main_df (조건부 서식 적용)
    - 매칭실패: unmatched_df
    """
    wb = Workbook()

    # === 1. 요약 시트 ===
    ws = wb.active
    ws.title = "요약"
    ws["A1"] = "페북성과보고서"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A2"] = period_label
    ws["A2"].font = Font(size=11, color="6B7280")

    if kpi:
        rows = [
            ("총 광고비", f"₩{kpi.get('total_spend', 0):,.0f}"),
            ("총 매출", f"₩{kpi.get('total_revenue', 0):,.0f}"),
            ("ROAS", f"{kpi.get('roas', 0):.0f}%"),
            ("전환수", f"{kpi.get('conversion_count', 0):,}"),
        ]
        for i, (label, value) in enumerate(rows, start=4):
            ws[f"A{i}"] = label
            ws[f"B{i}"] = value
            ws[f"A{i}"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 22

    # === 2. 광고별상세 시트 ===
    detail_ws = wb.create_sheet("광고별상세")
    _write_dataframe_with_format(detail_ws, main_df)

    # === 3. 매칭실패 시트 ===
    if unmatched_df is not None and not unmatched_df.empty:
        unmatched_ws = wb.create_sheet("매칭실패")
        _write_dataframe_with_format(unmatched_ws, unmatched_df, apply_pct_format=False)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_dataframe_with_format(
    ws,
    df: pd.DataFrame,
    apply_pct_format: bool = True,
) -> None:
    """DataFrame을 워크시트에 쓰면서 조건부 서식 적용."""
    if df.empty:
        ws["A1"] = "(데이터 없음)"
        return

    # 헤더
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=str(col_name))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    pct_cols_idx = {
        col_idx for col_idx, col_name in enumerate(df.columns, start=1)
        if col_name in PERCENT_COLUMNS and apply_pct_format
    }
    money_cols_idx = {
        col_idx for col_idx, col_name in enumerate(df.columns, start=1)
        if is_money_column(col_name)
    }
    validity_idx = None
    for col_idx, col_name in enumerate(df.columns, start=1):
        if col_name == "유효":
            validity_idx = col_idx
            break

    for row_idx, (_, row) in enumerate(df.iterrows(), start=2):
        is_invalid = validity_idx is not None and row.get("유효") == "X"
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=_safe_value(value))
            # 숫자 포맷
            if col_idx in money_cols_idx and isinstance(value, (int, float)) and not pd.isna(value):
                cell.number_format = '"₩"#,##0'
            elif col_idx in pct_cols_idx and isinstance(value, (int, float)) and not pd.isna(value):
                cell.number_format = '0"%"'

            if is_invalid:
                cell.fill = GRAY_FILL
                cell.font = GRAY_FONT
            elif col_idx in pct_cols_idx:
                try:
                    v = float(value)
                    if v >= 100:
                        cell.fill = PINK_FILL
                        cell.font = PINK_FONT
                    else:
                        cell.fill = BLUE_FILL
                        cell.font = BLUE_FONT
                except (TypeError, ValueError):
                    pass

    # 컬럼 너비 자동
    for col_idx, col_name in enumerate(df.columns, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, min(28, len(str(col_name)) * 2 + 4))


def _safe_value(value):
    """엑셀에 쓸 수 있는 값으로 변환."""
    if pd.isna(value):
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value
